import io
import json
from typing import Dict, Any

from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

import docx

class ExportService:
    @staticmethod
    def export_markdown(report_data: Dict[str, Any]) -> str:
        content = f"# {report_data.get('title', 'ResearchMind AI Report')}\n\n"
        content += f"**Date:** {report_data.get('date', 'N/A')}\n"
        content += f"**Fact-Check Confidence Score:** {report_data.get('confidence_score', 0) * 100:.1f}%\n\n"
        
        content += "## Executive Summary\n"
        content += f"{report_data.get('executive_summary', '')}\n\n"

        content += "## Key Findings\n"
        for finding in report_data.get('key_findings', []):
            content += f"- {finding}\n"
        content += "\n"

        if report_data.get('charts'):
            content += "## Data & Visualizations\n"
            for chart in report_data.get('charts', []):
                content += f"### {chart.get('title', 'Chart')}\n"
                content += f"Type: {chart.get('type', 'bar')}\n\n"
                content += "| Category | Value |\n|---|---|\n"
                for data_pt in chart.get('data', []):
                    cat = data_pt.get('label') or data_pt.get('category') or data_pt.get('name', '')
                    val = data_pt.get('value') or data_pt.get('val', '')
                    content += f"| {cat} | {val} |\n"
                content += "\n"

        content += "## Claim Verification & Conflicts\n"
        for claim in report_data.get('claims', []):
            content += f"- **Claim:** {claim.get('claim_text')}\n"
            content += f"  - Status: {claim.get('status')}\n"
            content += f"  - Confidence: {claim.get('confidence_score', 0) * 100:.1f}%\n"

        if report_data.get('conflicts'):
            content += "\n### Detected Evidence Conflicts\n"
            for conflict in report_data.get('conflicts', []):
                content += f"- **Conflict:** {conflict.get('reason')}\n"
                if conflict.get('resolution_note'):
                    content += f"  - *Resolution Note:* {conflict.get('resolution_note')}\n"

        content += "\n## Numbered References & Sources\n"
        for idx, src in enumerate(report_data.get('sources', []), start=1):
            content += f"[{idx}] **{src.get('title')}** - [{src.get('url')}]({src.get('url')}) (Publisher: {src.get('publisher', 'Unknown')})\n"

        if report_data.get('limitations'):
            content += f"\n## Limitations\n{report_data.get('limitations')}\n"

        return content

    @staticmethod
    def export_pdf(report_data: Dict[str, Any]) -> bytes:
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'DocTitle',
            parent=styles['Heading1'],
            fontSize=22,
            leading=26,
            textColor=colors.HexColor('#0F172A'),
            spaceAfter=12
        )
        h2_style = ParagraphStyle(
            'SectionH2',
            parent=styles['Heading2'],
            fontSize=15,
            leading=18,
            textColor=colors.HexColor('#1E293B'),
            spaceBefore=12,
            spaceAfter=6
        )
        body_style = ParagraphStyle(
            'DocBody',
            parent=styles['Normal'],
            fontSize=10,
            leading=14,
            textColor=colors.HexColor('#334155'),
            spaceAfter=6
        )

        elements = []

        # Title & Meta
        elements.append(Paragraph(report_data.get('title', 'ResearchMind AI Report'), title_style))
        meta_text = f"Confidence Score: <b>{report_data.get('confidence_score', 0) * 100:.1f}%</b> | Sources Analyzed: <b>{len(report_data.get('sources', []))}</b>"
        elements.append(Paragraph(meta_text, body_style))
        elements.append(Spacer(1, 10))

        # Executive Summary
        elements.append(Paragraph("Executive Summary", h2_style))
        elements.append(Paragraph(report_data.get('executive_summary', ''), body_style))
        elements.append(Spacer(1, 10))

        # Key Findings
        elements.append(Paragraph("Key Findings", h2_style))
        for finding in report_data.get('key_findings', []):
            elements.append(Paragraph(f"• {finding}", body_style))
        elements.append(Spacer(1, 10))

        # Data Tables if charts exist
        if report_data.get('charts'):
            elements.append(Paragraph("Extracted Data & Visualizations", h2_style))
            for chart in report_data.get('charts', []):
                elements.append(Paragraph(f"<b>{chart.get('title', 'Chart')}</b>", body_style))
                table_data = [["Metric / Category", "Value / Stat"]]
                for dp in chart.get('data', []):
                    cat = dp.get('label') or dp.get('category') or dp.get('name', '')
                    val = str(dp.get('value') or dp.get('val', ''))
                    table_data.append([cat, val])
                
                t = Table(table_data, colWidths=[250, 200])
                t.setStyle(TableStyle([
                    ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#E2E8F0')),
                    ('TEXTCOLOR', (0,0), (-1,0), colors.HexColor('#0F172A')),
                    ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
                    ('BOTTOMPADDING', (0,0), (-1,0), 6),
                    ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
                ]))
                elements.append(t)
                elements.append(Spacer(1, 10))

        # Sources
        elements.append(Paragraph("References & Cited Sources", h2_style))
        for idx, src in enumerate(report_data.get('sources', []), start=1):
            src_str = f"[{idx}] <b>{src.get('title')}</b> — {src.get('url')}"
            elements.append(Paragraph(src_str, body_style))

        doc.build(elements)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def export_docx(report_data: Dict[str, Any]) -> bytes:
        doc = docx.Document()
        
        # Title
        doc.add_heading(report_data.get('title', 'ResearchMind AI Report'), level=0)
        
        # Meta
        meta_p = doc.add_paragraph()
        meta_p.add_run(f"Confidence Score: {report_data.get('confidence_score', 0) * 100:.1f}%\n").bold = True
        meta_p.add_run(f"Total Sources: {len(report_data.get('sources', []))}")

        # Exec Summary
        doc.add_heading("Executive Summary", level=1)
        doc.add_paragraph(report_data.get('executive_summary', ''))

        # Key Findings
        doc.add_heading("Key Findings", level=1)
        for finding in report_data.get('key_findings', []):
            doc.add_paragraph(finding, style='List Bullet')

        # Data Tables
        if report_data.get('charts'):
            doc.add_heading("Extracted Statistics & Data", level=1)
            for chart in report_data.get('charts', []):
                doc.add_heading(chart.get('title', 'Chart Data'), level=2)
                table = doc.add_table(rows=1, cols=2)
                hdr_cells = table.rows[0].cells
                hdr_cells[0].text = 'Category / Parameter'
                hdr_cells[1].text = 'Extracted Metric'
                for dp in chart.get('data', []):
                    row_cells = table.add_row().cells
                    row_cells[0].text = str(dp.get('label') or dp.get('category') or dp.get('name', ''))
                    row_cells[1].text = str(dp.get('value') or dp.get('val', ''))

        # Sources
        doc.add_heading("Cited Sources", level=1)
        for idx, src in enumerate(report_data.get('sources', []), start=1):
            p = doc.add_paragraph(style='List Number')
            p.add_run(f"{src.get('title')} ").bold = True
            p.add_run(f"({src.get('url')})")

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
